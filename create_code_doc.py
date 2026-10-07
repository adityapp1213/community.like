from docx import Document
from docx.shared import Inches, Pt, RGBColor

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(.75)
section.bottom_margin = Inches(.75)
section.left_margin = Inches(.85)
section.right_margin = Inches(.85)
doc.styles['Normal'].font.name = 'Aptos'
doc.styles['Normal'].font.size = Pt(10)
doc.styles['Title'].font.name = 'Aptos Display'
doc.styles['Title'].font.size = Pt(24)
doc.styles['Title'].font.bold = True
doc.styles['Title'].font.color.rgb = RGBColor(38, 32, 52)
doc.styles['Heading 1'].font.name = 'Aptos Display'
doc.styles['Heading 1'].font.color.rgb = RGBColor(91, 62, 150)

doc.add_paragraph('Community Like Project Guide', style='Title')
doc.add_paragraph('A simple overview of the current student project community site.')
doc.add_heading('What the site does', level=1)
doc.add_paragraph('Community Like lets students share projects, ask for help, discover events, connect with friends, and manage their account. The main feed is focused on student work and university communities.')
doc.add_heading('Main files', level=1)
for name, description in [('index.html', 'Projects feed and project composer.'), ('event.html', 'Build sessions, peer help events, and project showcases.'), ('friends.html', 'Community member list and friend actions.'), ('profile.html', 'Profile information and projects.'), ('settings.html', 'Clerk account settings and logout.'), ('landing.html', 'The Clerk sign-in page.'), ('app.js', 'Renders pages, search, projects, likes, saves, friends, events, uploads, and deletion.'), ('data.js', 'Stores demo data and browser activity in localStorage.'), ('auth.js', 'Loads Clerk, protects main pages, and handles account access.'), ('styles.css', 'Shared layout and responsive styling.')]:
    p = doc.add_paragraph(style='List Bullet'); p.add_run(f'{name}: ').bold = True; p.add_run(description)
doc.add_heading('Authentication', level=1)
doc.add_paragraph('Clerk provides sign in, sign up, account settings, and logout. The landing page contains the Clerk sign-in component. Signed-out visitors are sent there before opening the main site.')
doc.add_heading('Current data storage', level=1)
doc.add_paragraph('The demo uses localStorage for projects, likes, saves, friends, events, and profile activity. This works in one browser but is not shared between users or devices.')
doc.add_heading('Next database step', level=1)
doc.add_paragraph('Supabase is the recommended shared database. Clerk user IDs can connect to Supabase profiles, projects, friends, followers, events, likes, and saves. The existing data layer can then use database calls instead of localStorage.')
doc.save('C:/Users/adity/OneDrive/Desktop/community.like/community-like-code-guide.docx')
